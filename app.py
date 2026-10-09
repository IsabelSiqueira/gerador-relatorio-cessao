import os
import io
import openpyxl
from openpyxl.styles import Font, Border, Side, PatternFill, Alignment
from openpyxl.utils import get_column_letter
import streamlit as st

# Configuração da página web
st.set_page_config(
    page_title="Gerador de Planilha de Cessão",
    page_icon="📊",
    layout="centered"
)

st.title("📊 Gerador de Planilha de Cessão")
st.write("Faça o upload do arquivo base e informe a marca para gerar a nova planilha formatada.")

# --- Formulário na Interface ---
uploaded_file = st.file_uploader("Selecione o arquivo Excel base (.xlsx)", type=["xlsx"])
marca_desejada = st.text_input("Nome da Marca (ex: ANIMALE)", "").strip()

btn_gerar = st.button("🚀 Gerar Planilha de Cessão", type="primary", use_container_width=True)

def processar_planilha(file_bytes, marca):
    SHEET_NOME = "Aposta e Venda"
    LINHA_CABECALHO_BASE = 6

    # Carregar workbook da memória
    wb_base = openpyxl.load_workbook(file_bytes, data_only=True)

    if SHEET_NOME not in wb_base.sheetnames:
        raise ValueError(f"A aba '{SHEET_NOME}' não foi encontrada no arquivo base.")

    ws_base = wb_base[SHEET_NOME]

    # Estrutura de colunas
    COLUNAS_BASE = [
        "CHAVE", "MARCA", "COLECAO_ETAPA", "PRODUTO", "COR", "DESCRIÇÃO", "PREÇO",
        "VENDA", "VO1", "VO2", "VO3", "VO4", "VO5", "VO6", "VO7", "VENDA LÍQUIDA $",
        "PROG VAREJO", "PV1", "PV2", "PV3", "PV4", "PV5", "PV6", "PV7",
        "FALTA", "F1", "F2", "F3", "F4", "F5", "F6", "F7", "FALTA $",
        "QTDE_CESSAO", "EX1", "EX2", "EX3", "EX4", "EX5", "EX6", "EX7",
        "FALTA FINAL", "FF1", "FF2", "FF3", "FF4", "FF5", "FF6", "FF7", "FALTA FINAL $"
    ]

    COLUNAS_CALCULADAS = [
        "CESSÃO FINAL",
        "CF1", "CF2", "CF3", "CF4", "CF5", "CF6", "CF7",
        "CESSÃO FINAL $",
        "LIBERAR CESSÃO?"
    ]

    TODAS_COLUNAS = COLUNAS_BASE + COLUNAS_CALCULADAS
    colunas_1_a_7 = {col for col in TODAS_COLUNAS if any(col.endswith(str(i)) for i in range(1, 8))}

    # Mapear cabeçalhos
    header_row_base = [cell.value for cell in ws_base[LINHA_CABECALHO_BASE]]
    col_map_base = {}

    for col_name in COLUNAS_BASE:
        if col_name in header_row_base:
            col_map_base[col_name] = header_row_base.index(col_name) + 1

    colunas_obrigatorias = ["MARCA", "COLECAO_ETAPA", "FALTA", "QTDE_CESSAO", "PROG VAREJO"]
    for col_req in colunas_obrigatorias:
        if col_req not in col_map_base:
            raise ValueError(f"A coluna '{col_req}' não foi encontrada na linha {LINHA_CABECALHO_BASE} da planilha base.")

    idx_marca = col_map_base["MARCA"]
    idx_colecao = col_map_base["COLECAO_ETAPA"]
    idx_falta = col_map_base["FALTA"]
    idx_qtde_cessao = col_map_base["QTDE_CESSAO"]
    idx_prog_varejo = col_map_base["PROG VAREJO"]

    def to_float(val):
        try:
            if val is None or val == "":
                return 0.0
            return float(val)
        except (ValueError, TypeError):
            return 0.0

    # Filtrar dados
    linhas_validas = []
    colecoes_encontradas = set()

    for row in ws_base.iter_rows(min_row=LINHA_CABECALHO_BASE + 1, values_only=True):
        val_marca = str(row[idx_marca - 1]).strip() if row[idx_marca - 1] is not None else ""
        
        if val_marca.upper() == marca.upper():
            falta_val = to_float(row[idx_falta - 1])
            qtde_cessao_val = to_float(row[idx_qtde_cessao - 1])
            prog_varejo_val = to_float(row[idx_prog_varejo - 1])
            
            if (falta_val > 0 or qtde_cessao_val > 0) and (prog_varejo_val > 0):
                linhas_validas.append(row)
                val_colecao = str(row[idx_colecao - 1]).strip() if row[idx_colecao - 1] is not None else ""
                if val_colecao:
                    colecoes_encontradas.add(val_colecao)

    if not linhas_validas:
        return None, f"Nenhum produto encontrado para a marca '{marca}' que atenda aos critérios.", ""

    colecao_str = "_".join(sorted(colecoes_encontradas)) if colecoes_encontradas else "GERAL"
    titulo_texto = f"CESSÃO {marca.upper()} {colecao_str.upper()}"

    # Criar nova planilha
    wb_nova = openpyxl.Workbook()
    ws_nova = wb_nova.active
    ws_nova.title = f"Cessão {marca}"
    ws_nova.views.sheetView[0].showGridLines = False

    # Título B2
    cell_titulo = ws_nova.cell(row=2, column=2, value=titulo_texto)
    cell_titulo.font = Font(size=18, bold=True)

    LINHA_INICIO_TABELA = 4
    COLUNA_INICIO_TABELA = 2

    # Estilos
    borda_cabecalho = Border(
        bottom=Side(style='thin', color='000000'),
        top=Side(style='thin', color='E0E0E0'),
        left=Side(style='thin', color='E0E0E0'),
        right=Side(style='thin', color='E0E0E0')
    )
    borda_dados_pontilhada = Border(
        bottom=Side(style='dotted', color='D3D3D3'),
        left=Side(style='dotted', color='D3D3D3'),
        right=Side(style='dotted', color='D3D3D3')
    )
    borda_instrucao = Border(
        top=Side(style='thin', color='D6B656'),
        bottom=Side(style='thin', color='D6B656'),
        left=Side(style='thin', color='D6B656'),
        right=Side(style='thin', color='D6B656')
    )

    fill_cinza_claro = PatternFill(start_color="F2F2F2", end_color="F2F2F2", fill_type="solid")
    fill_instrucao = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")

    col_letters = {}

    # Cabeçalhos
    for idx, col_name in enumerate(TODAS_COLUNAS):
        col_num = COLUNA_INICIO_TABELA + idx
        col_letter = get_column_letter(col_num)
        col_letters[col_name] = col_letter
        
        cell_head = ws_nova.cell(row=LINHA_INICIO_TABELA, column=col_num, value=col_name)
        cell_head.font = Font(bold=True)
        cell_head.border = borda_cabecalho

    # Caixa de instrução
    col_liberacao_idx = COLUNA_INICIO_TABELA + TODAS_COLUNAS.index("LIBERAR CESSÃO?")
    col_liberacao_letra = col_letters["LIBERAR CESSÃO?"]

    ws_nova.merge_cells(f"{col_liberacao_letra}2:{col_liberacao_letra}3")
    cell_instrucao = ws_nova.cell(row=2, column=col_liberacao_idx)
    cell_instrucao.value = (
        "📌 Como funciona a Cessão:\n"
        "• Atualmente está preenchida com os valores acordados no início da coleção (somente o necessário).\n"
        "• Digite 'OK' para liberar toda a cessão necessária e zerar a falta do produto.\n"
        "• Para ceder apenas uma parte, altere manualmente a quantidade diretamente na coluna por tamanho (EX1 a EX7)."
    )
    cell_instrucao.font = Font(size=9, italic=True)
    cell_instrucao.fill = fill_instrucao
    cell_instrucao.alignment = Alignment(wrap_text=True, vertical="center", horizontal="left")

    for r in range(2, 4):
        ws_nova.cell(row=r, column=col_liberacao_idx).border = borda_instrucao

    ws_nova.row_dimensions[2].height = 25
    ws_nova.row_dimensions[3].height = 55
    ws_nova.column_dimensions[col_liberacao_letra].width = 32

    # Dados e Fórmulas
    linha_atual_nova = LINHA_INICIO_TABELA + 1

    colunas_moeda = {"PREÇO", "VENDA LÍQUIDA $", "FALTA $", "FALTA FINAL $", "CESSÃO FINAL $"}
    colunas_texto = {"CHAVE", "MARCA", "SUBMARCA", "COLECAO_ETAPA", "PRODUTO", "COR", "DESCRIÇÃO", "LIBERAR CESSÃO?"}

    FORMATO_MOEDA = 'R$ #,##0.00'
    FORMATO_INTEIRO = '#,##0'

    for row in linhas_validas:
        dados_linha = []
        for col_name in COLUNAS_BASE:
            idx = col_map_base.get(col_name)
            if idx and idx <= len(row):
                val = row[idx - 1]
                dados_linha.append(val if val is not None else "")
            else:
                dados_linha.append("")
        
        for idx, valor in enumerate(dados_linha):
            col_num = COLUNA_INICIO_TABELA + idx
            ws_nova.cell(row=linha_atual_nova, column=col_num, value=valor)
        
        # Fórmulas
        for i in range(1, 8):
            col_f = col_letters[f"F{i}"]
            col_ex = col_letters[f"EX{i}"]
            col_ff = col_letters[f"FF{i}"]
            col_lib = col_letters["LIBERAR CESSÃO?"]
            
            formula_cf = f'=IF({col_lib}{linha_atual_nova}="OK", MIN(N({col_f}{linha_atual_nova}), N({col_ex}{linha_atual_nova})) + N({col_ff}{linha_atual_nova}), MIN(N({col_f}{linha_atual_nova}), N({col_ex}{linha_atual_nova})))'
            col_idx_cf = COLUNA_INICIO_TABELA + TODAS_COLUNAS.index(f"CF{i}")
            ws_nova.cell(row=linha_atual_nova, column=col_idx_cf, value=formula_cf)

        cf1_l = col_letters["CF1"]
        cf7_l = col_letters["CF7"]
        formula_cf_total = f'=SUM({cf1_l}{linha_atual_nova}:{cf7_l}{linha_atual_nova})'
        col_idx_cf_tot = COLUNA_INICIO_TABELA + TODAS_COLUNAS.index("CESSÃO FINAL")
        ws_nova.cell(row=linha_atual_nova, column=col_idx_cf_tot, value=formula_cf_total)

        col_cf_tot_l = col_letters["CESSÃO FINAL"]
        col_preco_l = col_letters["PREÇO"]
        formula_cf_total_rs = f'={col_cf_tot_l}{linha_atual_nova} * N({col_preco_l}{linha_atual_nova})'
        col_idx_cf_tot_rs = COLUNA_INICIO_TABELA + TODAS_COLUNAS.index("CESSÃO FINAL $")
        ws_nova.cell(row=linha_atual_nova, column=col_idx_cf_tot_rs, value=formula_cf_total_rs)

        col_idx_lib = COLUNA_INICIO_TABELA + TODAS_COLUNAS.index("LIBERAR CESSÃO?")
        ws_nova.cell(row=linha_atual_nova, column=col_idx_lib, value="")

        for idx, col_name in enumerate(TODAS_COLUNAS):
            col_num = COLUNA_INICIO_TABELA + idx
            cell_data = ws_nova.cell(row=linha_atual_nova, column=col_num)
            cell_data.border = borda_dados_pontilhada
            
            if col_name in colunas_1_a_7:
                cell_data.fill = fill_cinza_claro
                
            if col_name in colunas_moeda:
                cell_data.number_format = FORMATO_MOEDA
            elif col_name not in colunas_texto:
                cell_data.number_format = FORMATO_INTEIRO

        linha_atual_nova += 1

    # AutoFilter
    primeira_col_letra = get_column_letter(COLUNA_INICIO_TABELA)
    ultima_col_letra = get_column_letter(COLUNA_INICIO_TABELA + len(TODAS_COLUNAS) - 1)
    ultima_linha = linha_atual_nova - 1
    ws_nova.auto_filter.ref = f"{primeira_col_letra}{LINHA_INICIO_TABELA}:{ultima_col_letra}{ultima_linha}"

    # Salvar em buffer de memória
    output_buffer = io.BytesIO()
    wb_nova.save(output_buffer)
    output_buffer.seek(0)
    
    nome_saida = f"Cessao_{marca.replace(' ', '_')}_{colecao_str}.xlsx"
    return output_buffer, None, nome_saida

# --- Execução da Ação ---
if btn_gerar:
    if not uploaded_file:
        st.error("Por favor, faça o upload de um arquivo Excel base.")
    elif not marca_desejada:
        st.error("Por favor, informe o nome da marca.")
    else:
        with st.spinner("Processando dados e gerando nova planilha..."):
            buffer_excel, erro, nome_arquivo = processar_planilha(uploaded_file, marca_desejada)
            
            if erro:
                st.warning(erro)
            else:
                st.success(f"Planilha para **{marca_desejada.upper()}** gerada com sucesso!")
                st.download_button(
                    label="📥 Baixar Planilha Processada",
                    data=buffer_excel,
                    file_name=nome_arquivo,
                    mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    use_container_width=True
                )